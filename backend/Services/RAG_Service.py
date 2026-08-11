"""
RAG Service - Qdrant-backed ingestion and retrieval pipeline.

Replaces the earlier TF-IDF implementation with real vector embeddings
(Azure OpenAI) stored in Qdrant, giving semantic search instead of pure
keyword overlap. Public interface (search, ingest_file, remove_file,
list_documents, overview) is unchanged from the TF-IDF version, so
RAGSearchTool / RAGAgent / main.py require zero changes.
"""

import base64
import io
import logging
import os
import uuid

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from langchain_core.messages import HumanMessage
from PIL import Image
from qdrant_client import QdrantClient
from qdrant_client.http import models as qmodels

from llm import llm

logger = logging.getLogger("finance_workflow")

DOCUMENTS_DIR = os.getenv("DOCUMENTS_DIR", "./data/uploads")
SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}

QDRANT_URL = os.getenv("QDRANT_URL", "http://localhost:6333")
QDRANT_API_KEY = os.getenv(
    "QDRANT_API_KEY"
)  # None for local Docker, set for Qdrant Cloud
QDRANT_COLLECTION = os.getenv("QDRANT_COLLECTION", "finance_documents")

AZURE_OPENAI_EMBEDDING_DEPLOYMENT = os.getenv(
    "AZURE_OPENAI_EMBEDDING_DEPLOYMENT", "text-embedding-3-small"
)
EMBEDDING_DIM = int(
    os.getenv("EMBEDDING_DIM", "1536")
)  # 1536 for text-embedding-3-small / ada-002


def _analyze_image_with_vision(image_bytes: bytes) -> str:
    """Uses Azure OpenAI Vision (gpt-4o-mini) to describe charts, figures, tables, and images."""
    try:
        b64_str = base64.b64encode(image_bytes).decode("utf-8")
        vision_llm = llm.get_langchain_llm()
        message = HumanMessage(
            content=[
                {
                    "type": "text",
                    "text": (
                        "You are an expert financial document analyst. Analyze this image, chart, table, or figure in full detail. "
                        "Extract all numbers, data points, labels, visual trends, column headers, and financial insights clearly."
                    ),
                },
                {
                    "type": "image_url",
                    "image_url": {"url": f"data:image/png;base64,{b64_str}"},
                },
            ]
        )
        res = vision_llm.invoke([message])
        return res.content.strip()
    except Exception as e:
        logger.warning("Vision analysis failed for image chunk: %s", e)
        return ""


def _extract_multimodal_pdf(file_path: str) -> str:
    """Multimodal PDF parser using PyMuPDF and Azure OpenAI Vision.
    Extracts text, structured markdown tables, and AI vision summaries for visual figures & charts.
    """
    doc = fitz.open(file_path)
    full_sections = []

    for page_idx, page in enumerate(doc):
        page_num = page_idx + 1
        page_content = [f"--- Page {page_num} ---"]

        # 1. Extract structured text
        page_text = page.get_text("text") or ""

        # 2. Extract tables into clean Markdown tables
        try:
            tabs = page.find_tables()
            if tabs and tabs.tables:
                for t_idx, tab in enumerate(tabs.tables):
                    df_df = tab.extract()
                    if df_df and len(df_df) > 0:
                        headers = [str(c or "").strip() for c in df_df[0]]
                        md_rows = [
                            f"| {' | '.join(headers)} |",
                            f"| {' | '.join(['---'] * len(headers))} |",
                        ]
                        for row in df_df[1:]:
                            cells = [str(c or "").strip() for c in row]
                            md_rows.append(f"| {' | '.join(cells)} |")
                        md_table_str = "\n".join(md_rows)
                        page_content.append(
                            f"\n[TABLE (Page {page_num}, Table {t_idx + 1})]\n{md_table_str}\n"
                        )
        except Exception as te:
            logger.debug("Table extraction skipped on page %s: %s", page_num, te)

        # 3. Extract embedded images & charts for Vision AI analysis
        try:
            image_list = page.get_images(full=True)
            for img_idx, img_info in enumerate(image_list[:3]):
                xref = img_info[0]
                base_img = doc.extract_image(xref)
                if not base_img:
                    continue
                img_bytes = base_img["image"]

                try:
                    pil_img = Image.open(io.BytesIO(img_bytes))
                    w, h = pil_img.size
                    if w < 100 or h < 100:
                        continue
                except Exception:
                    pass

                vision_summary = _analyze_image_with_vision(img_bytes)
                if vision_summary:
                    page_content.append(
                        f"\n[VISUAL CHART / IMAGE ANALYSIS (Page {page_num}, Fig {img_idx + 1})]\n{vision_summary}\n"
                    )
        except Exception as ie:
            logger.debug("Image extraction skipped on page %s: %s", page_num, ie)

        # 4. Fallback for scanned PDF pages (low text, high graphics)
        if len(page_text.strip()) < 40:
            try:
                pix = page.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("png")
                page_vision = _analyze_image_with_vision(img_bytes)
                if page_vision:
                    page_content.append(
                        f"\n[SCANNED PAGE OCR ANALYSIS (Page {page_num})]\n{page_vision}\n"
                    )
            except Exception as pe:
                logger.debug("Page rendering OCR skipped on page %s: %s", page_num, pe)

        if page_text.strip():
            page_content.append(page_text.strip())

        full_sections.append("\n".join(page_content))

    doc.close()
    return "\n\n".join(full_sections)


def _extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    if ext == ".pdf":
        return _extract_multimodal_pdf(file_path)
    if ext == ".docx":
        doc = DocxDocument(file_path)
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"Unsupported file type: {ext}")


def _chunk_text(text: str, chunk_size: int = 400, overlap: int = 80) -> list[str]:
    words = text.split()
    if not words:
        return []
    chunks = []
    start = 0
    while start < len(words):
        end = start + chunk_size
        chunks.append(" ".join(words[start:end]))
        start += chunk_size - overlap
    return chunks


class EmbeddingClient:
    """Thin wrapper around Azure OpenAI embeddings, isolated so it's swappable/mockable."""

    def __init__(self):
        self._model = None

    def _get_model(self):
        if self._model is None:
            from langchain_openai import AzureOpenAIEmbeddings

            self._model = AzureOpenAIEmbeddings(
                azure_deployment=AZURE_OPENAI_EMBEDDING_DEPLOYMENT,
                api_version=os.getenv("AZURE_OPENAI_API_VERSION", "2024-02-01"),
                azure_endpoint=os.getenv("AZURE_OPENAI_ENDPOINT"),
                api_key=os.getenv("AZURE_OPENAI_API_KEY"),
            )
        return self._model

    def embed_query(self, text: str) -> list[float]:
        return self._get_model().embed_query(text)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return self._get_model().embed_documents(texts)


class RAGService:
    _instance = None

    def __new__(cls, *args, **kwargs):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
            cls._instance._initialized = False
        return cls._instance

    def __init__(
        self,
        client: QdrantClient | None = None,
        embedder: EmbeddingClient | None = None,
    ):
        if self._initialized:
            return
        self._initialized = True
        os.makedirs(DOCUMENTS_DIR, exist_ok=True)
        self._embedder = embedder or EmbeddingClient()
        if client:
            self._client = client
            self._ensure_collection()
        else:
            try:
                self._client = QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY)
                self._ensure_collection()
            except Exception as e:
                logger.warning(
                    f"Could not connect to Qdrant cloud at {QDRANT_URL}: {e}. Falling back to local Qdrant storage."
                )
                qdrant_path = os.getenv("QDRANT_PATH", "./data/qdrant_db")
                os.makedirs(qdrant_path, exist_ok=True)
                self._client = QdrantClient(path=qdrant_path)
                self._ensure_collection()

    def _ensure_collection(self):
        existing = [c.name for c in self._client.get_collections().collections]
        if QDRANT_COLLECTION not in existing:
            self._client.create_collection(
                collection_name=QDRANT_COLLECTION,
                vectors_config=qmodels.VectorParams(
                    size=EMBEDDING_DIM, distance=qmodels.Distance.COSINE
                ),
            )
            logger.info(
                "Created Qdrant collection '%s' (dim=%s)",
                QDRANT_COLLECTION,
                EMBEDDING_DIM,
            )
        try:
            self._client.create_payload_index(
                collection_name=QDRANT_COLLECTION,
                field_name="source",
                field_schema=qmodels.PayloadSchemaType.KEYWORD,
            )
        except Exception:
            pass

    def ingest_file(self, file_path: str) -> int:
        ext = os.path.splitext(file_path)[1].lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {ext}")

        filename = os.path.basename(file_path)
        text = _extract_text(file_path)
        chunks = _chunk_text(text)
        if not chunks:
            return 0

        # re-upload of the same filename replaces its previous points
        self._delete_points_for_source(filename)

        vectors = self._embedder.embed_documents(chunks)
        points = [
            qmodels.PointStruct(
                id=str(uuid.uuid4()),
                vector=vector,
                payload={"source": filename, "text": chunk},
            )
            for chunk, vector in zip(chunks, vectors)
        ]
        self._client.upsert(collection_name=QDRANT_COLLECTION, points=points)
        return len(points)

    def _delete_points_for_source(self, filename: str):
        self._client.delete(
            collection_name=QDRANT_COLLECTION,
            points_selector=qmodels.FilterSelector(
                filter=qmodels.Filter(
                    must=[
                        qmodels.FieldCondition(
                            key="source", match=qmodels.MatchValue(value=filename)
                        )
                    ]
                )
            ),
        )

    def sync_documents(self):
        """Ensures all supported files in DOCUMENTS_DIR are indexed into Qdrant."""
        if not os.path.exists(DOCUMENTS_DIR):
            return
        indexed_sources = set()
        try:
            next_offset = None
            while True:
                records, next_offset = self._client.scroll(
                    collection_name=QDRANT_COLLECTION,
                    with_payload=["source"],
                    limit=200,
                    offset=next_offset,
                )
                for record in records:
                    s = record.payload.get("source")
                    if s:
                        indexed_sources.add(s)
                if next_offset is None:
                    break
        except Exception:
            pass

        for fname in os.listdir(DOCUMENTS_DIR):
            ext = os.path.splitext(fname)[1].lower()
            if ext in SUPPORTED_EXTENSIONS and fname not in indexed_sources:
                fpath = os.path.join(DOCUMENTS_DIR, fname)
                if os.path.isfile(fpath):
                    try:
                        self.ingest_file(fpath)
                    except Exception as e:
                        logger.warning("Auto-sync failed for %s: %s", fname, e)

    def search(self, query: str, k: int = 3) -> list[dict]:
        try:
            count = self._client.count(collection_name=QDRANT_COLLECTION).count
        except Exception:
            count = 0

        if count == 0:
            self.sync_documents()
            try:
                count = self._client.count(collection_name=QDRANT_COLLECTION).count
            except Exception:
                count = 0

        if count == 0:
            return []

        query_vector = self._embedder.embed_query(query)
        response = self._client.query_points(
            collection_name=QDRANT_COLLECTION,
            query=query_vector,
            limit=k,
            with_payload=True,
        )
        return [
            {
                "source": point.payload.get("source", "unknown"),
                "text": point.payload.get("text", ""),
                "score": float(point.score),
            }
            for point in response.points
        ]

    def overview(self, chunks_per_doc: int = 2) -> list[dict]:
        """
        Fallback sample for meta-queries ('summarize the documents'). Less
        critical now that real embeddings handle these reasonably well via
        semantic similarity, but kept as a safety net.
        """
        seen_counts: dict[str, int] = {}
        results: list[dict] = []
        next_offset = None
        while True:
            records, next_offset = self._client.scroll(
                collection_name=QDRANT_COLLECTION,
                with_payload=True,
                limit=100,
                offset=next_offset,
            )
            for record in records:
                source = record.payload.get("source", "unknown")
                count = seen_counts.get(source, 0)
                if count >= chunks_per_doc:
                    continue
                seen_counts[source] = count + 1
                results.append(
                    {
                        "source": source,
                        "text": record.payload.get("text", ""),
                        "score": None,
                    }
                )
            if next_offset is None:
                break
        return results

    def remove_file(self, filename: str) -> bool:
        safe_name = os.path.basename(filename)  # prevent path traversal
        path = os.path.join(DOCUMENTS_DIR, safe_name)
        if not os.path.isfile(path):
            return False
        os.remove(path)
        self._delete_points_for_source(safe_name)
        return True

    def list_documents(self) -> list[str]:
        self.sync_documents()
        sources = set()
        if os.path.exists(DOCUMENTS_DIR):
            for fname in os.listdir(DOCUMENTS_DIR):
                if os.path.splitext(fname)[1].lower() in SUPPORTED_EXTENSIONS:
                    sources.add(fname)
        try:
            next_offset = None
            while True:
                records, next_offset = self._client.scroll(
                    collection_name=QDRANT_COLLECTION,
                    with_payload=["source"],
                    limit=200,
                    offset=next_offset,
                )
                for record in records:
                    s = record.payload.get("source")
                    if s:
                        sources.add(s)
                if next_offset is None:
                    break
        except Exception:
            pass
        return sorted(sources)


rag_service = RAGService()
